"""
Build a formatted Excel screening report with Pandas and OpenPyXL.

The export is for human review. It includes identifiers, job-related scores,
skills, experience, and evidence. It omits phone numbers and email addresses
by default so the spreadsheet does not spread extra contact data.
"""

from __future__ import annotations

from io import BytesIO
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.job_matcher import MatchResult


EXPORT_COLUMNS = [
    "Rank",
    "Resume file",
    "Candidate name",
    "Job-match score",
    "Required-skill coverage",
    "Preferred-skill coverage",
    "Experience overlap",
    "Education overlap",
    "Supporting-evidence score",
    "Matched skills",
    "Missing/unverified required skills",
    "Unverified preferred skills",
    "Relevant experience",
    "Evidence quotes",
    "Score explanation",
    "Warnings",
    "Human-review note",
]


def build_report_dataframe(matches: Iterable[MatchResult]) -> pd.DataFrame:
    rows = []
    for match in matches:
        rows.append(
            {
                "Rank": match.rank,
                "Resume file": match.original_filename,
                "Candidate name": match.display_name,
                "Job-match score": match.overall_score,
                "Required-skill coverage": match.component_scores.get("required_skills"),
                "Preferred-skill coverage": match.component_scores.get("preferred_skills"),
                "Experience overlap": match.component_scores.get("relevant_experience"),
                "Education overlap": match.component_scores.get("education"),
                "Supporting-evidence score": match.component_scores.get("supporting_evidence"),
                "Matched skills": _join(match.matched_skills),
                "Missing/unverified required skills": _join(match.missing_required_skills),
                "Unverified preferred skills": _join(match.unverified_preferred_skills),
                "Relevant experience": match.relevant_experience_summary,
                "Evidence quotes": _join(match.evidence_quotes, separator=" | "),
                "Score explanation": " ".join(match.score_explanation),
                "Warnings": _join(match.warnings, separator=" | "),
                "Human-review note": (
                    "Score is job-requirement overlap only. It is not a hiring decision "
                    "or a prediction of job performance. Review the evidence before any outcome."
                ),
            }
        )
    if not rows:
        return pd.DataFrame(columns=EXPORT_COLUMNS)
    return pd.DataFrame(rows, columns=EXPORT_COLUMNS)


def generate_excel_bytes(matches: Iterable[MatchResult]) -> bytes:
    """Return an .xlsx file in memory so Streamlit can offer a download."""
    frame = build_report_dataframe(matches)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Screening Results")
        worksheet = writer.book["Screening Results"]
        _format_worksheet(worksheet, frame)
        _add_readme_sheet(writer.book)
    return buffer.getvalue()


def _format_worksheet(worksheet: Worksheet, frame: pd.DataFrame) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    widths = {
        "A": 10,
        "B": 28,
        "C": 22,
        "D": 18,
        "E": 22,
        "F": 24,
        "G": 20,
        "H": 20,
        "I": 24,
        "J": 36,
        "K": 36,
        "L": 32,
        "M": 40,
        "N": 44,
        "O": 44,
        "P": 32,
        "Q": 40,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    score_columns = {"D", "E", "F", "G", "H", "I"}
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, max_col=worksheet.max_column):
        worksheet.row_dimensions[row[0].row].height = 48
        for cell in row:
            cell.alignment = wrap
            letter = get_column_letter(cell.column)
            if letter in score_columns and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0"

    worksheet.auto_filter.ref = f"A1:{get_column_letter(frame.shape[1])}{max(frame.shape[0] + 1, 1)}"


def _add_readme_sheet(workbook) -> None:
    sheet = workbook.create_sheet("How to read this report")
    lines = [
        "This workbook is a screening aid, not a hiring decision.",
        "The job-match score is a weighted overlap of job-related evidence from the resume.",
        "Candidates are not automatically rejected.",
        "Phone numbers and email addresses are omitted from this export.",
        "Do not interpret names or other identifiers as ranking features.",
        "Always read the evidence quotes and missing-skill lists before any human decision.",
    ]
    sheet["A1"] = "Report notes"
    sheet["A1"].font = Font(bold=True, size=14)
    for index, line in enumerate(lines, start=3):
        sheet[f"A{index}"] = line
        sheet[f"A{index}"].alignment = Alignment(wrap_text=True)
    sheet.column_dimensions["A"].width = 110


def _join(values: list[str] | None, separator: str = ", ") -> str:
    if not values:
        return ""
    return separator.join(str(item) for item in values if item)
