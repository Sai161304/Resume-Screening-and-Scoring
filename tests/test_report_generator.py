"""Excel export tests with fictional screening results."""

from openpyxl import load_workbook

from src.job_matcher import JobRequirements, match_resume_to_job
from src.report_generator import build_report_dataframe, generate_excel_bytes
from src.resume_parser import parsed_resume_from_dict


def test_excel_report_contains_scores_and_omits_contact_columns():
    parsed = parsed_resume_from_dict(
        {
            "candidate_name": "Alex Rivera",
            "email": "should-not-be-exported@example.test",
            "phone": "000-000-0000",
            "technical_skills": ["python", "sql"],
            "experience": [{"title": "Intern", "highlights": ["Used Python"]}],
        },
        "alex.pdf",
    )
    job = JobRequirements(
        job_title="Engineer",
        required_skills=["python", "sql"],
        preferred_skills=["aws"],
        required_qualifications=[],
        experience_themes=["backend"],
        source="test",
    )
    match = match_resume_to_job(parsed, job, "python sql intern backend")
    match.rank = 1

    frame = build_report_dataframe([match])
    assert "email" not in [column.lower() for column in frame.columns]
    assert "phone" not in [column.lower() for column in frame.columns]
    assert frame.loc[0, "Resume file"] == "alex.pdf"
    assert frame.loc[0, "Job-match score"] >= 0

    workbook_bytes = generate_excel_bytes([match])
    from io import BytesIO

    workbook = load_workbook(BytesIO(workbook_bytes))
    assert "Screening Results" in workbook.sheetnames
    sheet = workbook["Screening Results"]
    assert sheet.auto_filter.ref
    assert sheet.freeze_panes == "A2"
    headers = [cell.value for cell in sheet[1]]
    assert "Job-match score" in headers
    assert "Evidence quotes" in headers
